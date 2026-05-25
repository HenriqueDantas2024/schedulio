"""
Converte os 3 documentos Markdown para PDF (ReportLab) e DOCX (python-docx).
Uso: python docs/convert_docs.py
"""

import re
from pathlib import Path

# ─── config ─────────────────────────────────────────────────────────────────

DOCS_DIR = Path(__file__).parent

FILES = [
    ("manual-usuario.md",        "Manual do Usuário"),
    ("documentacao-tecnica.md",  "Documentação Técnica"),
    ("asbuilt-v2.md",            "As-Built v2"),
]

C_NAVY  = (26/255,  31/255,  54/255)
C_RED   = (232/255, 25/255,  60/255)
C_GRAY  = (100/255, 110/255, 130/255)
C_LIGHT = (0.95, 0.96, 0.98)
C_WHITE = (1, 1, 1)
C_BLACK = (0.12, 0.12, 0.16)

HEX_NAVY  = "1A1F36"
HEX_RED   = "E8193C"
HEX_LIGHT = "F5F6FA"
HEX_WHITE = "FFFFFF"

# ─── inline stripping ───────────────────────────────────────────────────────

def strip_md_inline(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*",     r"\1", text)
    text = re.sub(r"`(.+?)`",        r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text

def rl_escape(text: str) -> str:
    """Escape HTML special chars for ReportLab Paragraph."""
    return (text
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;"))

# ─── Markdown parser ─────────────────────────────────────────────────────────

def parse_lines(md: str):
    lines = md.splitlines()
    i = 0
    in_code = False
    code_buf = []

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if in_code:
                in_code = False
                yield "code_block", "\n".join(code_buf)
                code_buf = []
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.match(r"^[-: ]+$", c) for c in cells if c):
                i += 1
                continue
            is_header = (
                i + 1 < len(lines)
                and lines[i + 1].startswith("|")
                and all(re.match(r"^[-: ]+$", c.strip())
                        for c in lines[i + 1].strip("|").split("|") if c.strip())
            )
            yield ("table_header" if is_header else "table_row"), cells
            i += 1
            continue

        if   line.startswith("#### "): yield "heading3", line[5:].strip()
        elif line.startswith("### "):  yield "heading3", line[4:].strip()
        elif line.startswith("## "):   yield "heading2", line[3:].strip()
        elif line.startswith("# "):    yield "heading1", line[2:].strip()
        elif re.match(r"^[-*] ", line): yield "bullet",   line[2:].strip()
        elif re.match(r"^\d+\. ", line): yield "numbered", re.sub(r"^\d+\. ", "", line).strip()
        elif re.match(r"^---+$", line.strip()): yield "hr", ""
        elif line.strip() == "":        yield "blank", ""
        else:                           yield "paragraph", line.strip()
        i += 1


# ═══════════════════════════════════════════════════════════════════════════
# PDF via ReportLab
# ═══════════════════════════════════════════════════════════════════════════

def md_to_pdf(md_path: Path, pdf_path: Path, title: str):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, HRFlowable,
        Table, TableStyle, Preformatted, PageBreak,
    )
    from reportlab.lib.enums import TA_LEFT, TA_CENTER

    color_navy  = colors.Color(*C_NAVY)
    color_red   = colors.Color(*C_RED)
    color_gray  = colors.Color(*C_GRAY)
    color_light = colors.Color(*C_LIGHT)
    color_black = colors.Color(*C_BLACK)

    styles = getSampleStyleSheet()
    base = styles["Normal"]

    def S(name, parent=base, **kw):
        return ParagraphStyle(name, parent=parent, **kw)

    sH1 = S("H1", fontSize=17, textColor=color_navy, fontName="Helvetica-Bold",
             spaceAfter=4, spaceBefore=12, leading=22)
    sH2 = S("H2", fontSize=13, textColor=color_red, fontName="Helvetica-Bold",
             spaceAfter=4, spaceBefore=10, leading=18)
    sH3 = S("H3", fontSize=11, textColor=color_navy, fontName="Helvetica-Bold",
             spaceAfter=3, spaceBefore=8, leading=15)
    sP  = S("P",  fontSize=10, textColor=color_black, fontName="Helvetica",
             spaceAfter=4, leading=14)
    sBullet = S("Bullet", parent=sP, leftIndent=14, bulletIndent=4,
                spaceAfter=2, leading=13)
    sNum  = S("Num", parent=sP, leftIndent=14, spaceAfter=2, leading=13)
    sCode = S("Code", fontSize=8, fontName="Courier", textColor=color_navy,
              backColor=color_light, leftIndent=8, rightIndent=8,
              spaceBefore=4, spaceAfter=4, leading=11)
    sCover = S("Cover", fontSize=26, textColor=color_navy, fontName="Helvetica-Bold",
               leading=32)
    sSub   = S("Sub",   fontSize=11, textColor=color_gray, fontName="Helvetica",
               spaceAfter=16, leading=16)

    # --- page template with header/footer ---
    doc = SimpleDocTemplate(
        str(pdf_path), pagesize=A4,
        leftMargin=22*mm, rightMargin=22*mm,
        topMargin=28*mm, bottomMargin=22*mm,
    )

    def on_page(canvas, doc):
        w, h = A4
        # header bar
        canvas.setFillColor(color_navy)
        canvas.rect(0, h - 14*mm, w, 14*mm, fill=1, stroke=0)
        canvas.setFillColor(color_red)
        canvas.rect(0, h - 15.5*mm, w, 1.5*mm, fill=1, stroke=0)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.white)
        canvas.drawString(22*mm, h - 9*mm, f"IMP Concursos  |  {title}")
        canvas.drawRightString(w - 22*mm, h - 9*mm, f"Página {doc.page}")
        # footer
        canvas.setFillColor(color_red)
        canvas.rect(0, 10*mm, w, 1*mm, fill=1, stroke=0)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(color_gray)
        canvas.drawCentredString(w / 2, 6*mm, "grade-horaria-imp.vercel.app")

    def on_first_page(canvas, doc):
        w, h = A4
        canvas.setFillColor(color_navy)
        canvas.rect(0, h - 52*mm, w, 52*mm, fill=1, stroke=0)
        canvas.setFillColor(color_red)
        canvas.rect(0, h - 54*mm, w, 2*mm, fill=1, stroke=0)
        canvas.setFont("Helvetica-Bold", 11)
        canvas.setFillColor(colors.white)
        canvas.drawString(22*mm, h - 18*mm, "IMP Concursos")
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(colors.Color(1, 1, 1, 0.65))
        canvas.drawString(22*mm, h - 28*mm, "Grade Horária — Sistema de Gestão")
        # footer
        canvas.setFillColor(color_red)
        canvas.rect(0, 10*mm, w, 1*mm, fill=1, stroke=0)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(color_gray)
        canvas.drawCentredString(w / 2, 6*mm, "grade-horaria-imp.vercel.app")

    story = []
    # Cover section
    story.append(Spacer(1, 8*mm))
    story.append(Paragraph(title, sCover))
    story.append(Paragraph("Grade Horária IMP — 2026-05-20", sSub))
    story.append(HRFlowable(width="100%", thickness=1.5, color=color_red, spaceAfter=12))

    content = md_path.read_text(encoding="utf-8")
    table_headers: list = []
    table_rows: list    = []
    in_table = False
    num_counter = 0

    def flush_table():
        nonlocal table_headers, table_rows, in_table
        if not in_table:
            return
        all_data = []
        if table_headers:
            all_data.append([rl_escape(strip_md_inline(c)) for c in table_headers])
        for r in table_rows:
            all_data.append([rl_escape(strip_md_inline(str(c))) for c in r])
        if all_data:
            col_n = max(len(r) for r in all_data)
            # normalise row lengths
            for r in all_data:
                while len(r) < col_n:
                    r.append("")
            w_each = (doc.width) / col_n
            tbl = Table(all_data, colWidths=[w_each] * col_n, repeatRows=1 if table_headers else 0)
            ts = TableStyle([
                ("FONTNAME",    (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE",    (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.Color(*C_LIGHT), colors.white]),
                ("GRID",        (0, 0), (-1, -1), 0.3, colors.Color(0.85, 0.87, 0.92)),
                ("TOPPADDING",  (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ])
            if table_headers:
                ts.add("BACKGROUND", (0, 0), (-1, 0), color_navy)
                ts.add("TEXTCOLOR",  (0, 0), (-1, 0), colors.white)
                ts.add("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold")
            tbl.setStyle(ts)
            story.append(tbl)
            story.append(Spacer(1, 4))
        table_headers.clear()
        table_rows.clear()
        in_table = False

    for kind, text in parse_lines(content):
        if kind == "table_header":
            flush_table()
            in_table = True
            table_headers[:] = text
            continue
        elif kind == "table_row":
            in_table = True
            table_rows.append(text)
            continue
        else:
            flush_table()

        if kind == "heading1":
            num_counter = 0
            story.append(Paragraph(rl_escape(strip_md_inline(text)), sH1))
            story.append(HRFlowable(width="100%", thickness=0.5, color=color_navy, spaceAfter=4))
        elif kind == "heading2":
            num_counter = 0
            story.append(Paragraph(rl_escape(strip_md_inline(text)), sH2))
        elif kind == "heading3":
            num_counter = 0
            story.append(Paragraph(rl_escape(strip_md_inline(text)), sH3))
        elif kind == "paragraph":
            story.append(Paragraph(rl_escape(strip_md_inline(text)), sP))
        elif kind == "bullet":
            story.append(Paragraph(f"• {rl_escape(strip_md_inline(text))}", sBullet))
        elif kind == "numbered":
            num_counter += 1
            story.append(Paragraph(f"{num_counter}. {rl_escape(strip_md_inline(text))}", sNum))
        elif kind == "code_block":
            story.append(Preformatted(text, sCode))
        elif kind == "hr":
            story.append(HRFlowable(width="100%", thickness=0.5, color=color_gray,
                                    spaceBefore=6, spaceAfter=6))
        elif kind == "blank":
            story.append(Spacer(1, 2))

    flush_table()

    doc.build(story, onFirstPage=on_first_page, onLaterPages=on_page)
    print(f"  PDF  -> {pdf_path.name}")


# ═══════════════════════════════════════════════════════════════════════════
# DOCX via python-docx
# ═══════════════════════════════════════════════════════════════════════════

def md_to_docx(md_path: Path, docx_path: Path, title: str):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    RNAVY = RGBColor(26, 31, 54)
    RRED  = RGBColor(232, 25, 60)
    RGRAY = RGBColor(100, 110, 130)
    RWHT  = RGBColor(255, 255, 255)
    RBLK  = RGBColor(30, 30, 40)

    def cell_bg(cell, hex_color: str):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def para_bg(paragraph, hex_color: str):
        pPr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), hex_color)
        pPr.append(shd)

    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    # Cover
    p = doc.add_paragraph()
    r = p.add_run("IMP Concursos")
    r.font.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGRAY

    h = doc.add_paragraph()
    r = h.add_run(title)
    r.font.bold = True; r.font.size = Pt(26); r.font.color.rgb = RNAVY

    s = doc.add_paragraph()
    r = s.add_run("Grade Horária IMP — 2026-05-20")
    r.font.size = Pt(11); r.font.color.rgb = RGRAY

    doc.add_paragraph()

    content = md_path.read_text(encoding="utf-8")
    table_headers: list = []
    table_rows: list    = []
    in_table = False
    num_counter = 0

    def flush_table():
        nonlocal table_headers, table_rows, in_table
        if not in_table:
            return
        all_data = ([table_headers] if table_headers else []) + table_rows
        if not all_data:
            in_table = False; return
        col_n = max(len(r) for r in all_data)
        tbl = doc.add_table(rows=0, cols=col_n)
        tbl.style = "Table Grid"
        if table_headers:
            hrow = tbl.add_row()
            for j, ct in enumerate(table_headers[:col_n]):
                cell = hrow.cells[j]
                cell.text = strip_md_inline(ct)
                cell_bg(cell, HEX_NAVY)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                        run.font.color.rgb = RWHT
                        run.font.size = Pt(9)
        for i, row_data in enumerate(table_rows):
            drow = tbl.add_row()
            bg = HEX_LIGHT if i % 2 == 0 else HEX_WHITE
            for j, ct in enumerate(row_data[:col_n]):
                cell = drow.cells[j]
                cell.text = strip_md_inline(str(ct))
                cell_bg(cell, bg)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
        doc.add_paragraph()
        table_headers.clear(); table_rows.clear()
        in_table = False

    for kind, text in parse_lines(content):
        if kind == "table_header":
            flush_table(); in_table = True; table_headers[:] = text; continue
        elif kind == "table_row":
            in_table = True; table_rows.append(text); continue
        else:
            flush_table()

        if kind == "heading1":
            num_counter = 0
            p = doc.add_paragraph()
            r = p.add_run(strip_md_inline(text))
            r.font.bold = True; r.font.size = Pt(16); r.font.color.rgb = RNAVY
            # underline via border
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bot = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "8"); bot.set(qn("w:color"), HEX_NAVY)
            pBdr.append(bot); pPr.append(pBdr)

        elif kind == "heading2":
            num_counter = 0
            p = doc.add_paragraph()
            r = p.add_run(strip_md_inline(text))
            r.font.bold = True; r.font.size = Pt(13); r.font.color.rgb = RRED

        elif kind == "heading3":
            num_counter = 0
            p = doc.add_paragraph()
            r = p.add_run(strip_md_inline(text))
            r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = RNAVY

        elif kind == "paragraph":
            p = doc.add_paragraph()
            r = p.add_run(strip_md_inline(text))
            r.font.size = Pt(10); r.font.color.rgb = RBLK

        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(strip_md_inline(text))
            r.font.size = Pt(10)

        elif kind == "numbered":
            num_counter += 1
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(strip_md_inline(text))
            r.font.size = Pt(10)

        elif kind == "code_block":
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.font.name = "Courier New"; r.font.size = Pt(8); r.font.color.rgb = RNAVY
            para_bg(p, HEX_LIGHT)

        elif kind == "hr":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bot = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6"); bot.set(qn("w:color"), HEX_RED)
            pBdr.append(bot); pPr.append(pBdr)

        elif kind == "blank":
            pass

    flush_table()
    doc.save(str(docx_path))
    print(f"  DOCX -> {docx_path.name}")


# ─── main ───────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    out = DOCS_DIR / "output"
    out.mkdir(exist_ok=True)

    for filename, title in FILES:
        md_path = DOCS_DIR / filename
        if not md_path.exists():
            print(f"  SKIP: {filename}")
            continue
        stem = md_path.stem
        print(f"\n[{title}]")
        md_to_pdf(md_path,  out / f"{stem}.pdf",  title)
        md_to_docx(md_path, out / f"{stem}.docx", title)

    print(f"\nPronto! Arquivos em {DOCS_DIR / 'output'}")
